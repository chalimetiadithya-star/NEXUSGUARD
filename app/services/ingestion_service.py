import io
import re
from typing import List, Optional
from sqlalchemy.orm import Session
from app.models.document import Document, DocumentChunk
from app.models.enums import ClearanceLevel, DocumentStatus
from app.schemas.document import DocumentCreate

class IngestionService:
    """
    Handles file text extraction, security metadata validation,
    document record creation, and chunk indexing.
    """

    @staticmethod
    def extract_text_from_file(filename: str, file_bytes: bytes) -> str:
        ext = filename.lower().split(".")[-1] if "." in filename else ""

        if ext == "txt" or ext == "md":
            try:
                return file_bytes.decode("utf-8")
            except UnicodeDecodeError:
                return file_bytes.decode("latin-1")

        elif ext == "pdf":
            try:
                import pypdf
                reader = pypdf.PdfReader(io.BytesIO(file_bytes))
                text_parts = [page.extract_text() or "" for page in reader.pages]
                return "\n".join(text_parts).strip()
            except Exception as e:
                raise ValueError(f"Failed to parse PDF document: {str(e)}")

        elif ext == "docx":
            try:
                import docx
                doc = docx.Document(io.BytesIO(file_bytes))
                paragraphs = [p.text for p in doc.paragraphs if p.text]
                return "\n".join(paragraphs).strip()
            except Exception as e:
                raise ValueError(f"Failed to parse DOCX document: {str(e)}")

        else:
            # Attempt plain text read as fallback
            try:
                return file_bytes.decode("utf-8")
            except Exception:
                raise ValueError(f"Unsupported file extension: .{ext}. Supported formats: TXT, PDF, DOCX")

    @classmethod
    def validate_metadata(cls, data: DocumentCreate) -> None:
        if not data.document_id or not data.document_id.strip():
            raise ValueError("Document ID is required")
        if not data.title or not data.title.strip():
            raise ValueError("Document title is required")
        if not data.department or not data.department.strip():
            raise ValueError("Department is required")
        if not data.lineage_id or not data.lineage_id.strip():
            raise ValueError("Lineage ID is required")
        if not data.effective_date or not data.effective_date.strip():
            raise ValueError("Effective date is required")

        # Validate classification enum
        clearance = ClearanceLevel.from_string(data.classification)
        if not clearance:
            raise ValueError(f"Invalid classification '{data.classification}'. Must be one of: Public, Internal, Confidential, Restricted")

    @classmethod
    def chunk_content(cls, content: str, chunk_size: int = 500, overlap: int = 50) -> List[str]:
        words = content.split()
        if not words:
            return []
        chunks = []
        step = chunk_size - overlap if chunk_size > overlap else chunk_size
        for i in range(0, len(words), step):
            chunk_words = words[i:i + chunk_size]
            chunks.append(" ".join(chunk_words))
        return chunks

    @classmethod
    def ingest_document(cls, db: Session, doc_data: DocumentCreate) -> Document:
        # Validate security metadata before saving
        cls.validate_metadata(doc_data)

        # Check for duplicate document_id
        existing = db.query(Document).filter(Document.document_id == doc_data.document_id).first()
        if existing:
            # Update existing document
            existing.title = doc_data.title
            existing.owner = doc_data.owner
            existing.department = doc_data.department
            existing.classification = doc_data.classification
            existing.allowed_departments = doc_data.allowed_departments
            existing.allowed_roles = doc_data.allowed_roles
            existing.allowed_users = doc_data.allowed_users
            existing.denied_users = doc_data.denied_users
            existing.lineage_id = doc_data.lineage_id
            existing.version = doc_data.version
            existing.effective_date = doc_data.effective_date
            existing.status = doc_data.status
            existing.content = doc_data.content

            # Delete old chunks
            db.query(DocumentChunk).filter(DocumentChunk.document_id == existing.document_id).delete()
            db_doc = existing
        else:
            db_doc = Document(
                document_id=doc_data.document_id,
                title=doc_data.title,
                owner=doc_data.owner,
                department=doc_data.department,
                classification=doc_data.classification,
                lineage_id=doc_data.lineage_id,
                version=doc_data.version,
                effective_date=doc_data.effective_date,
                status=doc_data.status,
                content=doc_data.content
            )
            db_doc.allowed_departments = doc_data.allowed_departments
            db_doc.allowed_roles = doc_data.allowed_roles
            db_doc.allowed_users = doc_data.allowed_users
            db_doc.denied_users = doc_data.denied_users
            db.add(db_doc)

        db.flush()

        # Create chunks and search index
        chunks_text = cls.chunk_content(doc_data.content)
        for idx, c_text in enumerate(chunks_text):
            chunk = DocumentChunk(
                document_id=db_doc.document_id,
                chunk_index=idx,
                content=c_text,
                search_text=c_text.lower()
            )
            db.add(chunk)

        db.commit()
        db.refresh(db_doc)
        return db_doc
